import http.server
import urllib.request
import json
import os
import time
import base64
import secrets
import urllib.parse
import sqlite3
import random
import subprocess
import zipfile
import io
import shutil
import threading
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False
try:
    from cryptography.fernet import Fernet, InvalidToken
    HAS_CRYPTOGRAPHY = True
except ImportError:
    Fernet = None
    InvalidToken = Exception
    HAS_CRYPTOGRAPHY = False

PORT = 8001
GITHUB_REPO = "Hootsworth/BDR_Prototype"
GITHUB_COMMITS_API = f"https://api.github.com/repos/{GITHUB_REPO}/commits/main"
GITHUB_ZIPBALL_URL = f"https://github.com/{GITHUB_REPO}/archive/refs/heads/main.zip"

def safe_urlopen(request, timeout=30):
    try:
        return urllib.request.urlopen(request, timeout=timeout)
    except urllib.error.URLError as ex:
        if "CERTIFICATE_VERIFY_FAILED" in str(ex):
            import ssl
            ctx = ssl._create_unverified_context()
            return urllib.request.urlopen(request, timeout=timeout, context=ctx)
        raise

def load_local_env():
    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.env')
    if not os.path.exists(env_path):
        return
    with open(env_path, encoding='utf-8') as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith('#') or '=' not in line:
                continue
            key, value = line.split('=', 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))

load_local_env()

GOOGLE_AUTH_URL = 'https://accounts.google.com/o/oauth2/v2/auth'
GOOGLE_TOKEN_URL = 'https://oauth2.googleapis.com/token'
GOOGLE_SCOPES = [
    'openid', 'email', 'profile',
    'https://www.googleapis.com/auth/gmail.modify',
    'https://www.googleapis.com/auth/gmail.send',
    'https://www.googleapis.com/auth/gmail.compose',
    'https://www.googleapis.com/auth/gmail.readonly',
    'https://www.googleapis.com/auth/calendar',
    'https://www.googleapis.com/auth/calendar.events',
    'https://www.googleapis.com/auth/calendar.freebusy',
]
google_sessions = {}
google_oauth_states = {}
DATA_DIR = os.environ.get('PROTOTYPE_DATA_DIR', '/tmp/gtm-data' if os.environ.get('VERCEL') else os.path.join(os.path.dirname(os.path.abspath(__file__)), '.prototype-data'))
DB_PATH = os.path.join(DATA_DIR, 'gtm.sqlite3')
WORKBOOK_PATH = os.environ.get('GTM_WORKBOOK_PATH', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'gtm-console-database.xlsx'))
WORKBOOK_SHEETS = [
    'Contacts', 'Companies', 'Enrichment', 'Campaigns', 'Activities',
    'Approvals', 'Events', 'Settings', 'Runs', 'Metadata'
]
workbook_lock = threading.Lock()

def token_cipher():
    key = os.environ.get('TOKEN_ENCRYPTION_KEY')
    if not key or not HAS_CRYPTOGRAPHY or not Fernet:
        return None
    return Fernet(key.encode('utf-8'))

def encrypt_token(value):
    if not value:
        return None
    cipher = token_cipher()
    if cipher:
        return cipher.encrypt(value.encode('utf-8')).decode('ascii')
    return base64.b64encode(value.encode('utf-8')).decode('ascii')

def decrypt_token(value):
    if not value:
        return None
    cipher = token_cipher()
    if cipher:
        try:
            return cipher.decrypt(value.encode('ascii')).decode('utf-8')
        except Exception as ex:
            raise RuntimeError('Stored Google token cannot be decrypted; reconnect Google Workspace.') from ex
    try:
        return base64.b64decode(value.encode('ascii')).decode('utf-8')
    except Exception:
        return value

def db():
    os.makedirs(DATA_DIR, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute('CREATE TABLE IF NOT EXISTS app_state (key TEXT PRIMARY KEY, value TEXT NOT NULL, updated_at TEXT NOT NULL)')
    connection.execute('CREATE TABLE IF NOT EXISTS google_connections (id INTEGER PRIMARY KEY CHECK (id = 1), session_id TEXT, email TEXT, name TEXT, access_token TEXT, refresh_token TEXT, expires_at REAL, updated_at TEXT NOT NULL)')
    columns = {row['name'] for row in connection.execute('PRAGMA table_info(google_connections)').fetchall()}
    if 'session_id' not in columns:
        connection.execute('ALTER TABLE google_connections ADD COLUMN session_id TEXT')
    connection.execute('CREATE TABLE IF NOT EXISTS sent_emails (id INTEGER PRIMARY KEY AUTOINCREMENT, fingerprint TEXT UNIQUE NOT NULL, recipient TEXT NOT NULL, subject TEXT NOT NULL, provider_id TEXT, sent_at TEXT NOT NULL)')
    connection.commit()
    return connection

def persist_state(key, value):
    with db() as connection:
        connection.execute('INSERT INTO app_state(key, value, updated_at) VALUES(?, ?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at', (key, json.dumps(value), time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())))

def read_state(key):
    with db() as connection:
        row = connection.execute('SELECT value FROM app_state WHERE key = ?', (key,)).fetchone()
        return json.loads(row['value']) if row else None

def default_workbook_state():
    return {
        'contacts': [], 'events': {}, 'stats': {'emailsSent': 0, 'linkedinSent': 0, 'callsMade': 0, 'enrichedCount': 0},
        'meetings': [], 'approvals': [], 'workflowRuns': [], 'currentOutboundSubtab': 'prospects', 'autoEnrich': False
    }

def _maybe_json(value):
    if isinstance(value, str):
        text = value.strip()
        if (text.startswith('{') and text.endswith('}')) or (text.startswith('[') and text.endswith(']')):
            try:
                return json.loads(text)
            except ValueError:
                return value
    return value

def sheet_records(worksheet):
    if worksheet is None:
        return []
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration:
        return []
    records = []
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue
        record = {}
        for key, value in zip(headers, row):
            if key is None:
                continue
            record[key] = _maybe_json('' if value is None else value)
        records.append(record)
    return records

def contact_for_workbook_record(contacts_by_id, contacts_by_email, record):
    contact_id = record.get('contactId')
    if contact_id is not None and str(contact_id) in contacts_by_id:
        return contacts_by_id[str(contact_id)]
    email = record.get('email')
    if email:
        return contacts_by_email.get(email)
    return None

def read_workbook_state():
    json_path = WORKBOOK_PATH.replace('.xlsx', '.json')
    if not os.path.exists(WORKBOOK_PATH):
        if os.path.exists(json_path):
            try:
                with open(json_path, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                pass
        return default_workbook_state()

    if not HAS_OPENPYXL or openpyxl is None:
        if os.path.exists(json_path):
            try:
                with open(json_path, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                pass
        return default_workbook_state()

    with workbook_lock:
        workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

    sheets = {name: (workbook[name] if name in workbook.sheetnames else None) for name in WORKBOOK_SHEETS}
    contacts = sheet_records(sheets['Contacts'])
    contacts_by_id = {str(c['id']): c for c in contacts if c.get('id') is not None}
    contacts_by_email = {c['email']: c for c in contacts if c.get('email')}

    for enrichment in sheet_records(sheets['Enrichment']):
        contact = contact_for_workbook_record(contacts_by_id, contacts_by_email, enrichment)
        if not contact:
            continue
        contact.update({
            'enriched': enrichment.get('enriched'),
            'enrichmentStatus': enrichment.get('enrichmentStatus'),
            'enrichmentSources': enrichment.get('enrichmentSources'),
            'enrichmentFields': enrichment.get('enrichmentFields'),
            'aiEnrichment': enrichment.get('aiEnrichment'),
            'enrichedAt': enrichment.get('enrichedAt')
        })

    for campaign in sheet_records(sheets['Campaigns']):
        contact = contact_for_workbook_record(contacts_by_id, contacts_by_email, campaign)
        if contact:
            contact.update(campaign)

    for activity in sheet_records(sheets['Activities']):
        contact = contact_for_workbook_record(contacts_by_id, contacts_by_email, activity)
        if contact and activity.get('type') == 'gmail_reply':
            contact.setdefault('replyHistory', [])
            if not any(r.get('messageId') == activity.get('messageId') for r in contact['replyHistory']):
                contact['replyHistory'].append(activity)

    events = {}
    for row in sheet_records(sheets['Events']):
        event_name = row.pop('eventName', None)
        if not event_name:
            continue
        events.setdefault(event_name, []).append(row)

    settings_values = {}
    for row in sheet_records(sheets['Settings']):
        key = row.get('key')
        if not key:
            continue
        raw_value = row.get('value')
        try:
            settings_values[key] = json.loads(raw_value) if isinstance(raw_value, str) else raw_value
        except ValueError:
            settings_values[key] = raw_value

    state = default_workbook_state()
    state['contacts'] = contacts
    state['approvals'] = sheet_records(sheets['Approvals'])
    state['workflowRuns'] = sheet_records(sheets['Runs'])
    if events:
        state['events'] = events
    if 'stats' in settings_values:
        state['stats'] = settings_values['stats']
    if 'meetings' in settings_values:
        state['meetings'] = settings_values['meetings']
    if 'currentOutboundSubtab' in settings_values:
        state['currentOutboundSubtab'] = settings_values['currentOutboundSubtab']
    if 'autoEnrich' in settings_values:
        state['autoEnrich'] = bool(settings_values['autoEnrich'])
    return state

def workbook_rows_for(records):
    rows = []
    for record in records or []:
        row = {}
        for key, value in record.items():
            row[key] = json.dumps(value) if isinstance(value, (dict, list)) else value
        rows.append(row)
    return rows

def build_workbook_snapshot(state):
    contacts = state.get('contacts') or []

    companies_seen = {}
    for contact in contacts:
        company = contact.get('company')
        if not company or company in companies_seen:
            continue
        companies_seen[company] = {
            'company': company,
            'industry': contact.get('industry', ''),
            'contacts': sum(1 for c in contacts if c.get('company') == company)
        }

    enrichment = [{
        'contactId': c.get('id'), 'email': c.get('email'),
        'enriched': bool(c.get('enriched')), 'enrichmentStatus': c.get('enrichmentStatus', ''),
        'enrichmentSources': c.get('enrichmentSources') or [], 'enrichmentFields': c.get('enrichmentFields') or [],
        'aiEnrichment': c.get('aiEnrichment') or {}, 'enrichedAt': c.get('enrichedAt', '')
    } for c in contacts if c.get('aiEnrichment') or c.get('enrichmentStatus') or c.get('enriched')]

    campaigns = [{
        'contactId': c.get('id'), 'email': c.get('email'),
        'emailDraft': c.get('emailDraft') or {}, 'emailsSent': bool(c.get('emailsSent')),
        'emailSentAt': c.get('emailSentAt', ''), 'emailProviderId': c.get('emailProviderId', ''),
        'linkedinDraft': c.get('linkedinDraft'), 'linkedinSent': bool(c.get('linkedinSent')),
        'callsMade': c.get('callsMade') or []
    } for c in contacts if c.get('emailDraft') or c.get('emailsSent') or c.get('linkedinDraft') or c.get('linkedinSent') or c.get('callsMade')]

    activities = []
    for c in contacts:
        for reply in (c.get('replyHistory') or []):
            activities.append({'contactId': c.get('id'), 'email': c.get('email'), 'type': 'gmail_reply', **reply})

    events_rows = []
    for event_name, attendees in (state.get('events') or {}).items():
        for attendee in (attendees or []):
            events_rows.append({'eventName': event_name, **attendee})

    settings_rows = [
        {'key': 'events', 'value': json.dumps(state.get('events') or {})},
        {'key': 'stats', 'value': json.dumps(state.get('stats') or {})},
        {'key': 'meetings', 'value': json.dumps(state.get('meetings') or [])},
        {'key': 'currentOutboundSubtab', 'value': json.dumps(state.get('currentOutboundSubtab') or 'prospects')},
        {'key': 'autoEnrich', 'value': json.dumps(bool(state.get('autoEnrich')))},
        {'key': 'savedAt', 'value': json.dumps(time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()))}
    ]

    return {
        'Contacts': contacts,
        'Companies': list(companies_seen.values()),
        'Enrichment': enrichment,
        'Campaigns': campaigns,
        'Activities': activities,
        'Approvals': state.get('approvals') or [],
        'Events': events_rows,
        'Settings': settings_rows,
        'Runs': state.get('workflowRuns') or [],
        'Metadata': [{'schema': 'gtm-console-workbook-v2', 'exportedAt': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())}]
    }

def write_workbook_state(state):
    json_path = WORKBOOK_PATH.replace('.xlsx', '.json')
    if not HAS_OPENPYXL or openpyxl is None:
        with workbook_lock:
            os.makedirs(os.path.dirname(WORKBOOK_PATH) or '.', exist_ok=True)
            tmp_json = json_path + '.tmp'
            with open(tmp_json, 'w', encoding='utf-8') as f:
                json.dump(state, f, indent=2)
            os.replace(tmp_json, json_path)
        return

    snapshot = build_workbook_snapshot(state)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for name in WORKBOOK_SHEETS:
        sheet = workbook.create_sheet(title=name)
        rows = workbook_rows_for(snapshot.get(name) or [])
        headers = []
        for row in rows:
            for key in row.keys():
                if key not in headers:
                    headers.append(key)
        if headers:
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(header, '') for header in headers])

    with workbook_lock:
        os.makedirs(os.path.dirname(WORKBOOK_PATH) or '.', exist_ok=True)
        tmp_path = WORKBOOK_PATH + '.tmp'
        workbook.save(tmp_path)
        os.replace(tmp_path, WORKBOOK_PATH)
        # Also write JSON snapshot for quick reference and backup
        try:
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump(state, f, indent=2)
        except Exception:
            pass


def persist_google_connection(session):
    with db() as connection:
        connection.execute(
            'INSERT INTO google_connections(id, session_id, email, name, access_token, refresh_token, expires_at, updated_at) VALUES(1, ?, ?, ?, ?, ?, ?, ?) '
            'ON CONFLICT(id) DO UPDATE SET session_id=excluded.session_id, email=excluded.email, name=excluded.name, access_token=excluded.access_token, refresh_token=COALESCE(excluded.refresh_token, google_connections.refresh_token), expires_at=excluded.expires_at, updated_at=excluded.updated_at',
            (session.get('session_id'), session.get('email'), session.get('name'), encrypt_token(session.get('access_token')), encrypt_token(session.get('refresh_token')), session.get('expires_at'), time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()))
        )

def restore_google_session(session_id):
    with db() as connection:
        row = connection.execute('SELECT * FROM google_connections WHERE session_id = ?', (session_id,)).fetchone()
        if not row:
            return None
        restored = dict(row)
        restored['access_token'] = decrypt_token(restored.get('access_token'))
        restored['refresh_token'] = decrypt_token(restored.get('refresh_token'))
        return restored

def refresh_google_session(session):
    if not session.get('refresh_token') or session.get('expires_at', 0) > time.time() + 60:
        return session
    _, tokens = post_form(GOOGLE_TOKEN_URL, {
        'client_id': os.environ.get('GOOGLE_CLIENT_ID'),
        'client_secret': os.environ.get('GOOGLE_CLIENT_SECRET'),
        'refresh_token': session['refresh_token'], 'grant_type': 'refresh_token'
    })
    session['access_token'] = tokens['access_token']
    session['expires_at'] = time.time() + tokens.get('expires_in', 3600)
    persist_google_connection(session)
    google_sessions[session['session_id']] = session
    return session

def email_fingerprint(recipient, subject, body):
    import hashlib
    return hashlib.sha256(f'{recipient}\0{subject}\0{body}'.encode('utf-8')).hexdigest()

def check_send_guardrails(payload):
    recipient, subject, body = payload.get('to'), payload.get('subject'), payload.get('body')
    if not recipient or not subject or not body:
        return 'Recipient, subject, and body are required.'
    if payload.get('suppressed') or payload.get('unsubscribed'):
        return 'Recipient is suppressed or unsubscribed.'
    if payload.get('approved') is not True:
        return 'An explicit approval is required before sending.'
    fingerprint = email_fingerprint(recipient or '', subject or '', body or '')
    with db() as connection:
        if connection.execute('SELECT 1 FROM sent_emails WHERE fingerprint = ?', (fingerprint,)).fetchone():
            return 'This exact email has already been sent.'
        today = time.strftime('%Y-%m-%d', time.gmtime())
        daily_limit = int(os.environ.get('PROTOTYPE_DAILY_SEND_LIMIT', '25'))
        sent_today = connection.execute('SELECT COUNT(*) AS count FROM sent_emails WHERE sent_at LIKE ?', (f'{today}%',)).fetchone()['count']
        if sent_today >= daily_limit:
            return f'Daily prototype send limit ({daily_limit}) reached.'
    return None

def record_sent_email(payload, provider_id):
    with db() as connection:
        connection.execute('INSERT OR IGNORE INTO sent_emails(fingerprint, recipient, subject, provider_id, sent_at) VALUES(?, ?, ?, ?, ?)', (email_fingerprint(payload['to'], payload['subject'], payload['body']), payload['to'], payload['subject'], provider_id, time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())))

def gmail_replies(session, payload):
    contacts = payload.get('contacts', [])[:50]
    replies = []
    for contact in contacts:
        email = contact.get('email') if isinstance(contact, dict) else contact
        if not email:
            continue
        query = urllib.parse.quote(f'from:{email} newer_than:30d')
        _, listing = google_api_get(f'https://gmail.googleapis.com/gmail/v1/users/me/messages?q={query}&maxResults=10', session['access_token'])
        for message in listing.get('messages', []):
            _, detail = google_api_get(f"https://gmail.googleapis.com/gmail/v1/users/me/messages/{message['id']}?format=metadata&metadataHeaders=From&metadataHeaders=Subject&metadataHeaders=Date", session['access_token'])
            headers = {h['name'].lower(): h['value'] for h in detail.get('payload', {}).get('headers', [])}
            replies.append({'contactEmail': email, 'messageId': detail.get('id'), 'threadId': detail.get('threadId'), 'from': headers.get('from', email), 'subject': headers.get('subject', ''), 'date': headers.get('date', ''), 'snippet': detail.get('snippet', '')})
    return {'replies': replies, 'syncedAt': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())}

def json_response(handler, status, payload):
    body = json.dumps(payload).encode('utf-8')
    handler.send_response(status)
    handler.send_header('Content-Type', 'application/json')
    handler.send_header('Cache-Control', 'no-store')
    handler.send_header('Content-Length', str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)

def post_json(url, payload, headers):
    return request_with_retry(url, json.dumps(payload).encode('utf-8'), {**headers, 'Content-Type': 'application/json'}, 'POST')

def post_form(url, payload):
    return request_with_retry(url, urllib.parse.urlencode(payload).encode('utf-8'), {'Content-Type': 'application/x-www-form-urlencoded'}, 'POST')

def request_with_retry(url, data, headers, method, attempts=3):
    last_error = None
    for attempt in range(attempts):
        try:
            request = urllib.request.Request(url, data=data, headers=headers, method=method)
            with safe_urlopen(request, timeout=30) as response:
                raw = response.read().decode('utf-8')
                return response.status, json.loads(raw) if raw else {}
        except urllib.error.HTTPError as ex:
            last_error = ex
            if ex.code not in (408, 429, 500, 502, 503, 504) or attempt == attempts - 1:
                raise
        except (urllib.error.URLError, TimeoutError) as ex:
            last_error = ex
            if attempt == attempts - 1:
                raise
        time.sleep((2 ** attempt) + random.random())
    raise last_error

def google_api(method, url, access_token, payload=None):
    headers = {'Authorization': f'Bearer {access_token}'}
    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode('utf-8') if payload is not None else None,
        headers={**headers, 'Content-Type': 'application/json'},
        method=method,
    )
    return request_with_retry(url, request.data, request.headers, method)

def google_api_get(url, access_token):
    request = urllib.request.Request(url, headers={'Authorization': f'Bearer {access_token}'}, method='GET')
    return request_with_retry(url, None, request.headers, 'GET')

def gmail_raw_message(to, subject, body):
    mime = f'To: {to}\r\nSubject: {subject}\r\nContent-Type: text/plain; charset="UTF-8"\r\n\r\n{body}'
    return base64.urlsafe_b64encode(mime.encode('utf-8')).decode('ascii').rstrip('=')

def gmail_request(session, resource, payload, send=False):
    raw = gmail_raw_message(payload['to'], payload['subject'], payload['body'])
    if send:
        url = 'https://gmail.googleapis.com/gmail/v1/users/me/messages/send'
        request_body = {'raw': raw}
    else:
        url = 'https://gmail.googleapis.com/gmail/v1/users/me/drafts'
        request_body = {'message': {'raw': raw}}
    return google_api('POST', url, session['access_token'], request_body)[1]

def get_local_version_info():
    version_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'version.json')
    local_sha = ""
    version_name = "1.0.0"
    if os.path.exists(version_file):
        try:
            with open(version_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                local_sha = data.get('commit', '')
                version_name = data.get('version', '1.0.0')
        except Exception:
            pass
    if not local_sha:
        try:
            local_sha = subprocess.check_output(
                ['git', 'rev-parse', 'HEAD'],
                cwd=os.path.dirname(os.path.abspath(__file__)),
                stderr=subprocess.DEVNULL
            ).decode('utf-8').strip()
        except Exception:
            local_sha = "unknown"
    return {"sha": local_sha, "version": version_name}

def is_same_commit(sha1, sha2):
    if not sha1 or not sha2 or sha1 == "unknown" or sha2 == "unknown":
        return False
    s1, s2 = str(sha1).strip().lower(), str(sha2).strip().lower()
    return s1 == s2 or s1.startswith(s2) or s2.startswith(s1)

def check_for_updates():
    local_info = get_local_version_info()
    local_sha = local_info.get("sha", "")
    req = urllib.request.Request(
        GITHUB_COMMITS_API,
        headers={"User-Agent": "GTM-Console-App", "Accept": "application/vnd.github.v3+json"}
    )
    try:
        with safe_urlopen(req, timeout=10) as response:
            if response.status == 200:
                data = json.loads(response.read().decode('utf-8'))
                remote_sha = data.get("sha", "")
                commit_info = data.get("commit", {})
                commit_msg = commit_info.get("message", "").split("\n")[0]
                commit_author = commit_info.get("author", {}).get("name", "")
                commit_date = commit_info.get("author", {}).get("date", "")
                
                is_update_available = bool(
                    remote_sha and local_sha and local_sha != "unknown" and not is_same_commit(local_sha, remote_sha)
                )
                
                return {
                    "update_available": is_update_available,
                    "current_version": local_info.get("version", "1.1.0"),
                    "current_commit": local_sha[:7] if local_sha != "unknown" else "v1.1.0",
                    "latest_commit": remote_sha[:7] if remote_sha else "",
                    "commit_message": commit_msg,
                    "author": commit_author,
                    "date": commit_date,
                    "repo_url": f"https://github.com/{GITHUB_REPO}"
                }
    except Exception as ex:
        return {
            "update_available": False,
            "error": str(ex),
            "current_version": local_info.get("version", "1.1.0"),
            "current_commit": local_info.get("sha", "")[:7]
        }
    return {
        "update_available": False,
        "current_version": local_info.get("version", "1.1.0"),
        "current_commit": local_info.get("sha", "")[:7]
    }

def apply_system_update():
    app_root = os.path.dirname(os.path.abspath(__file__))
    has_git = os.path.isdir(os.path.join(app_root, '.git'))
    
    if has_git:
        try:
            subprocess.check_call(['git', 'pull', 'origin', 'main'], cwd=app_root, timeout=60)
            new_sha = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=app_root).decode('utf-8').strip()
            
            version_file = os.path.join(app_root, 'version.json')
            with open(version_file, 'w', encoding='utf-8') as f:
                json.dump({
                    "version": "1.0.0",
                    "commit": new_sha,
                    "updated_at": time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
                }, f, indent=2)
                
            return {"status": "success", "mode": "git", "new_version": new_sha[:7]}
        except Exception:
            pass  # Fall through to zipball extraction if git fails
            
    req = urllib.request.Request(GITHUB_ZIPBALL_URL, headers={"User-Agent": "GTM-Console-App"})
    with safe_urlopen(req, timeout=45) as response:
        zip_bytes = response.read()
        
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
        if not names:
            raise RuntimeError("Empty archive received from GitHub.")
        prefix = names[0].split('/')[0] + '/'
        
        preserve_paths = {'.env', '.prototype-data', '.venv', 'node_modules', 'gtm.sqlite3', 'version.json'}
        
        for member in zf.infolist():
            if not member.filename.startswith(prefix):
                continue
            rel_path = member.filename[len(prefix):]
            if not rel_path or rel_path.endswith('/'):
                continue
                
            first_segment = rel_path.split('/')[0]
            if first_segment in preserve_paths or rel_path.endswith('.xlsx') or rel_path.endswith('.sqlite3'):
                continue
                
            target_file = os.path.join(app_root, rel_path)
            os.makedirs(os.path.dirname(target_file), exist_ok=True)
            with zf.open(member) as src, open(target_file, 'wb') as dst:
                shutil.copyfileobj(src, dst)
                
    # Fetch latest full commit sha from GitHub to write into version.json
    latest_full_sha = ""
    try:
        req_commit = urllib.request.Request(GITHUB_COMMITS_API, headers={"User-Agent": "GTM-Console-App", "Accept": "application/vnd.github.v3+json"})
        with safe_urlopen(req_commit, timeout=10) as c_resp:
            c_data = json.loads(c_resp.read().decode('utf-8'))
            latest_full_sha = c_data.get("sha", "")
    except Exception:
        pass

    latest_sha = latest_full_sha or "latest"
    version_file = os.path.join(app_root, 'version.json')
    with open(version_file, 'w', encoding='utf-8') as f:
        json.dump({
            "version": "1.1.0",
            "commit": latest_sha,
            "updated_at": time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
        }, f, indent=2)
        
    return {"status": "success", "mode": "zipball", "new_version": latest_sha[:7]}

class ProxyHTTPRequestHandler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        # Same-origin is the normal path; restrict cross-origin calls to local development.
        origin = self.headers.get('Origin')
        if origin in ('http://localhost:8001', 'http://127.0.0.1:8001'):
            self.send_header('Access-Control-Allow-Origin', origin)
            self.send_header('Access-Control-Allow-Credentials', 'true')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')
        self.send_header('Cache-Control', 'no-cache, no-store, must-revalidate')
        self.send_header('Pragma', 'no-cache')
        self.send_header('Expires', '0')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def _session_id(self):
        cookie = self.headers.get('Cookie', '')
        for part in cookie.split(';'):
            key, _, value = part.strip().partition('=')
            if key == 'gtm_session':
                return value
        return None

    def _google_session(self):
        session_id = self._session_id()
        if not session_id:
            return None
        session = google_sessions.get(session_id) or restore_google_session(session_id)
        if not session:
            return None
        try:
            return refresh_google_session(session)
        except Exception:
            return None

    def do_GET(self):
        if self.path == '/api/google/oauth/start':
            client_id = os.environ.get('GOOGLE_CLIENT_ID')
            redirect_uri = os.environ.get('GOOGLE_REDIRECT_URI', 'http://localhost:8001/api/google/oauth/callback')
            if not client_id:
                json_response(self, 503, {'error': 'GOOGLE_CLIENT_ID is not configured on the server.'})
                return
            state = secrets.token_urlsafe(32)
            google_oauth_states[state] = time.time()
            params = urllib.parse.urlencode({
                'client_id': client_id, 'redirect_uri': redirect_uri,
                'response_type': 'code', 'scope': ' '.join(GOOGLE_SCOPES),
                'access_type': 'offline', 'prompt': 'consent', 'state': state,
            })
            self.send_response(302)
            self.send_header('Location', f'{GOOGLE_AUTH_URL}?{params}')
            self.end_headers()
            return

        if self.path.startswith('/api/google/oauth/callback'):
            query = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
            state = query.get('state', [''])[0]
            code = query.get('code', [''])[0]
            if not state or state not in google_oauth_states or time.time() - google_oauth_states.pop(state) > 600:
                json_response(self, 400, {'error': 'Invalid or expired Google OAuth state.'})
                return
            if not code:
                json_response(self, 400, {'error': query.get('error', ['Google authorization was cancelled.'])[0]})
                return
            try:
                token_status, tokens = post_form(
                    GOOGLE_TOKEN_URL,
                    {'code': code, 'client_id': os.environ.get('GOOGLE_CLIENT_ID'),
                     'client_secret': os.environ.get('GOOGLE_CLIENT_SECRET'),
                     'redirect_uri': os.environ.get('GOOGLE_REDIRECT_URI', 'http://localhost:8001/api/google/oauth/callback'),
                     'grant_type': 'authorization_code'},
                )
                access_token = tokens['access_token']
                profile_status, profile = google_api('GET', 'https://openidconnect.googleapis.com/v1/userinfo', access_token)
                session_id = secrets.token_urlsafe(32)
                google_sessions[session_id] = {
                    'session_id': session_id,
                    'access_token': access_token, 'refresh_token': tokens.get('refresh_token'),
                    'expires_at': time.time() + tokens.get('expires_in', 3600),
                    'email': profile.get('email'), 'name': profile.get('name'),
                }
                persist_google_connection(google_sessions[session_id])
                self.send_response(302)
                self.send_header('Set-Cookie', f'gtm_session={session_id}; HttpOnly; SameSite=Lax; Path=/')
                self.send_header('Location', '/?google=connected')
                self.end_headers()
            except Exception as ex:
                json_response(self, 502, {'error': f'Google OAuth exchange failed: {ex}'})
            return

        if self.path == '/api/google/status':
            session = self._google_session()
            json_response(self, 200, {'connected': bool(session), 'email': session.get('email') if session else None})
            return

        if self.path == '/api/google/verify':
            session = self._google_session()
            if not session:
                json_response(self, 401, {'error': 'Connect a Google Workspace account first.'})
                return
            try:
                _, gmail_profile = google_api_get('https://gmail.googleapis.com/gmail/v1/users/me/profile', session['access_token'])
                _, calendar = google_api_get('https://www.googleapis.com/calendar/v3/calendars/primary', session['access_token'])
                json_response(self, 200, {'gmail': {'email': gmail_profile.get('emailAddress'), 'messagesTotal': gmail_profile.get('messagesTotal')}, 'calendar': {'id': calendar.get('id'), 'summary': calendar.get('summary')}, 'status': 'verified'})
            except urllib.error.HTTPError as ex:
                json_response(self, ex.code, {'error': ex.read().decode('utf-8', errors='replace')})
            return

        if self.path == '/api/google/disconnect':
            session_id = self._session_id()
            google_sessions.pop(session_id or '', None)
            with db() as connection:
                connection.execute('DELETE FROM google_connections WHERE session_id = ?', (session_id,))
            self.send_response(302)
            self.send_header('Set-Cookie', 'gtm_session=; Max-Age=0; Path=/')
            self.send_header('Location', '/')
            self.end_headers()
            return

        if self.path == '/api/state':
            json_response(self, 200, {'state': read_state('database') or {}})
            return

        if self.path == '/api/workbook/state':
            try:
                json_response(self, 200, {'state': read_workbook_state(), 'path': WORKBOOK_PATH})
            except Exception as ex:
                json_response(self, 500, {'error': f'Could not read the workbook: {ex}'})
            return

        if self.path == '/api/system/update-check':
            json_response(self, 200, check_for_updates())
            return

        super().do_GET()

    def do_POST(self):
        content_length = int(self.headers.get('Content-Length', 0))
        post_data = self.rfile.read(content_length)
        try:
            payload = json.loads(post_data.decode('utf-8')) if post_data else {}
        except json.JSONDecodeError:
            json_response(self, 400, {'error': 'Request body must be valid JSON.'})
            return

        if self.path == '/api/system/update':
            try:
                result = apply_system_update()
                json_response(self, 200, result)
            except Exception as ex:
                json_response(self, 500, {'error': f'System update failed: {ex}'})
            return

        if self.path.startswith('/api/google/'):
            session = self._google_session()
            if not session:
                json_response(self, 401, {'error': 'Connect a Google Workspace account first.'})
                return
            try:
                if self.path == '/api/google/gmail/draft':
                    result = gmail_request(session, 'drafts', payload, send=False)
                elif self.path == '/api/google/gmail/send':
                    guardrail_error = check_send_guardrails(payload)
                    if guardrail_error:
                        json_response(self, 409, {'error': guardrail_error})
                        return
                    result = gmail_request(session, 'messages/send', payload, send=True)
                    record_sent_email(payload, result.get('id'))
                elif self.path == '/api/google/calendar/freebusy':
                    result = google_api('POST', 'https://www.googleapis.com/calendar/v3/freeBusy', session['access_token'], payload)[1]
                elif self.path == '/api/google/calendar/events':
                    result = google_api('POST', 'https://www.googleapis.com/calendar/v3/calendars/primary/events?conferenceDataVersion=1', session['access_token'], payload)[1]
                elif self.path == '/api/google/gmail/replies':
                    result = gmail_replies(session, payload)
                else:
                    json_response(self, 404, {'error': 'Unknown Google API route.'})
                    return
                json_response(self, 200, result)
            except urllib.error.HTTPError as ex:
                detail = ex.read().decode('utf-8', errors='replace')
                json_response(self, ex.code, {'error': detail})
            except Exception as ex:
                json_response(self, 502, {'error': f'Google API request failed: {ex}'})
            return

        if self.path == '/api/state':
            persist_state('database', payload.get('state', {}))
            json_response(self, 200, {'status': 'saved'})
            return

        if self.path == '/api/workbook/state':
            try:
                write_workbook_state(payload.get('state', {}))
                json_response(self, 200, {'status': 'saved', 'savedAt': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())})
            except Exception as ex:
                json_response(self, 500, {'error': f'Could not save the workbook: {ex}'})
            return

        # AI enrichment stays server-side so provider credentials never reach the browser.
        if self.path == '/api/ai/enrich':
            api_key = os.environ.get('OPENAI_API_KEY')
            if not api_key:
                json_response(self, 503, {'error': 'OPENAI_API_KEY is not configured on the server.'})
                return
            contacts = payload.get('contacts', [])
            if not isinstance(contacts, list) or not contacts:
                json_response(self, 400, {'error': 'At least one contact is required.'})
                return
            prompt = {
                'contacts': contacts[:10],
                'requested_fields': payload.get('fields') or ['professional_summary', 'seniority', 'department', 'likely_pain_points', 'buying_signals', 'personalization_angles', 'relevant_topics', 'data_gaps', 'confidence'],
                'instructions': (
                    'Return one profile per input contact using requested_fields. Use only supplied facts '
                    'and cautious inferences. Never invent phone numbers, emails, URLs, employers, events, or '
                    'specific claims. Mark inferred values and leave unknown fields empty. Return JSON only.'
                )
            }
            try:
                _, result = post_json(
                    'https://api.openai.com/v1/chat/completions',
                    {'model': 'gpt-4o-mini', 'temperature': 0.2, 'response_format': {'type': 'json_object'},
                     'messages': [{'role': 'system', 'content': 'You are a careful B2B research assistant.'},
                                  {'role': 'user', 'content': json.dumps(prompt)}]},
                    {'Authorization': f'Bearer {api_key}'}
                )
                content = result['choices'][0]['message']['content']
                parsed = json.loads(content)
                profiles = parsed.get('profiles', parsed.get('results', parsed if isinstance(parsed, list) else []))
                json_response(self, 200, {'profiles': profiles, 'provider': 'openai'})
            except Exception as ex:
                json_response(self, 502, {'error': f'AI enrichment failed: {ex}'})
            return

        # Transactional outbound over HTTPS; this uses Resend, not SMTP.
        if self.path == '/api/email/send':
            api_key = os.environ.get('RESEND_API_KEY')
            sender = os.environ.get('RESEND_FROM_EMAIL')
            recipient = payload.get('to')
            subject = payload.get('subject')
            body = payload.get('body')
            if not api_key or not sender:
                json_response(self, 503, {'error': 'RESEND_API_KEY and RESEND_FROM_EMAIL must be configured.'})
                return
            if not recipient or not subject or not body:
                json_response(self, 400, {'error': 'to, subject, and body are required.'})
                return
            try:
                status, result = post_json(
                    'https://api.resend.com/emails',
                    {'from': sender, 'to': [recipient], 'subject': subject, 'text': body},
                    {'Authorization': f'Bearer {api_key}'}
                )
                json_response(self, status, {'provider': 'resend', 'id': result.get('id'), 'status': 'sent'})
            except urllib.error.HTTPError as ex:
                detail = ex.read().decode('utf-8', errors='replace')
                json_response(self, ex.code, {'error': f'Email provider rejected the request: {detail}'})
            except Exception as ex:
                json_response(self, 502, {'error': f'Email send failed: {ex}'})
            return

        # Lemlist live sending is intentionally unavailable until a real provider
        # connection and server-side credential handling are implemented.
        if self.path == '/api/lemlist/send-test':
            json_response(self, 501, {'error': 'Lemlist live sending is not configured in this build.'})

        # Intercept and proxy requests destined for Explorium API
        elif self.path.startswith('/api/proxy/'):
            target_path = self.path[len('/api/proxy/'):]
            target_url = f"https://api.explorium.ai/{target_path}"
            
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            
            headers = {}
            for k, v in self.headers.items():
                if k.lower() not in ('host', 'api_key'):
                    headers[k] = v
            headers['Host'] = 'api.explorium.ai'
            if os.environ.get('EXPLORIUM_API_KEY'):
                headers['api_key'] = os.environ['EXPLORIUM_API_KEY']
            
            req = urllib.request.Request(target_url, data=post_data, headers=headers, method='POST')
            try:
                with safe_urlopen(req, timeout=30) as response:
                    res_data = response.read()
                    self.send_response(response.status)
                    self.send_header('Content-Type', response.headers.get('Content-Type', 'application/json'))
                    self.end_headers()
                    self.wfile.write(res_data)
            except urllib.error.HTTPError as e:
                res_data = e.read()
                self.send_response(e.code)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(res_data)
            except Exception as e:
                self.send_response(500)
                self.end_headers()
                self.wfile.write(json.dumps({'error': str(e)}).encode())
        else:
            super().do_POST()

if __name__ == '__main__':
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    print(f"Starting Local Campaign Console static server with API proxy on port {PORT}...")
    http.server.HTTPServer.allow_reuse_address = True
    try:
        server = http.server.HTTPServer(('0.0.0.0', PORT), ProxyHTTPRequestHandler)
    except OSError as err:
        if err.errno == 48:
            print(f"\n[ERROR] Port {PORT} is already in use by another running process.")
            print(f"Run 'lsof -ti:{PORT} | xargs kill -9' in terminal to terminate the previous instance and run again.\n")
        raise
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nShutting down server.")
